from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/mattdimock/Documents/Jobs/Job Search")
OUT_DIR = ROOT / "applications" / "hospitality-lanes"
OUT_FILE = OUT_DIR / "Matt Dimock - Broadway Barback Master - Resume.docx"

NAVY = RGBColor(24, 53, 77)
ACCENT = RGBColor(182, 111, 55)
TEXT = RGBColor(28, 31, 36)
MUTED = RGBColor(95, 104, 113)

SUMMARY = (
    "Reliable, team-first operator pursuing barback work in downtown Nashville. Brings hands-on retail leadership in stocking, "
    "restocking, inventory counts, opening and closing, customer support, and physical floor work, plus years of operating experience "
    "built around clean handoffs, calm execution, and doing the unglamorous work that keeps a busy team moving. Ready to earn trust "
    "quickly in a high-volume bar, entertainment, or restaurant environment."
)

READINESS = [
    "Immediate availability for nights, weekends, holidays, and high-volume downtown shifts",
    "Comfortable with long shifts, constant movement, lifting, stocking, and support work",
    "Strong follow-through, clear communication, and low-ego team support",
    "Fast learner with strong memory for standards, routines, floor layouts, and manager preferences",
    "Professional with guests and steady under pressure",
    "Ready to complete TABC-certified alcohol awareness training immediately",
]

EXPERIENCE = [
    ("Get Fractional", "Growth Systems Consultant / Owner", "Oct 2022-Present", []),
    ("Prosper Wireless", "Director of Growth & Retention", "Sept 2023-Nov 2025", []),
    ("Dealer Acceleration Group", "Co-Founder / Chief Revenue Officer", "Jan 2022-Dec 2022", []),
    ("Affordable Insurance Quotes", "Chief Marketing Officer", "Jun 2020-Dec 2021", []),
    ("Breakthrough Academy", "Marketing Operations Manager", "Aug 2019-Jun 2021", []),
    ("SkyFineUSA", "Chief Marketing Officer / Head of Growth", "Jan 2019-Aug 2019", []),
    ("Bob's Watches", "Director of Marketing", "Jan 2017-Dec 2018", []),
    ("Swell Marketing", "SEO Manager / SEO Director", "Feb 2016-Jun 2016", []),
    ("iMarket Solutions", "SEO Manager", "May 2013-Jan 2016", []),
    ("National Positions", "SEO Specialist / Director of Special Projects", "Sep 2007-May 2013", []),
    (
        "Lowe's Home Improvement",
        "Team Leader, Paint Department",
        "Mar 2005-Feb 2007",
        [
            "Promoted into a team-lead role within six months and supported customer service, merchandising, stocking, inventory counts, opening and closing, and day-to-day department execution for a four-person team in a fast-paced retail environment.",
        ],
    ),
    (
        "Boething Treeland Nursery",
        "Guide",
        "Dec 2003-Aug 2004",
        [
            "Promoted from Loader to Guide within three months, supporting customers, loading purchases, restocking product areas, maintaining grounds, and keeping daily operations moving in a physical retail environment.",
        ],
    ),
]

BEST_FIT = [
    "Broadway barback, bar support, venue support, restaurant support, entertainment hospitality",
    "Strong fit for teams that need reliability, stamina, coachability, and no-drama support during rush periods",
]


def set_margins(section):
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)


def set_bottom_border(paragraph, color="B66F37", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def set_font(run, size, color=TEXT, bold=False, name="Aptos"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_header(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("Matt Dimock"), 21, NAVY, True, "Aptos Display")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("Barback | Broadway Bar Support"), 10.5, ACCENT, True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(
        p.add_run("Mount Juliet, TN  |  805-620-2826  |  mattdim805@gmail.com  |  linkedin.com/in/mattdimock/"),
        8.7,
        MUTED,
    )
    set_bottom_border(p)


def add_heading(doc, text, before=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(text.upper()), 8.8, NAVY, True)


def add_paragraph(doc, text, size=8.8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(text), size)


def add_bullets(doc, items, size=8.3):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(9)
        p.paragraph_format.first_line_indent = Pt(-7)
        set_font(p.add_run(f"- {item}"), size)


def add_experience(doc, company, title, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1.8)
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(company), 8.8, TEXT, True)
    set_font(p.add_run(f"  |  {title}  |  {dates}"), 8.05, MUTED)
    add_bullets(doc, bullets, size=8.15)


def build():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    set_margins(section)

    add_header(doc)
    add_heading(doc, "Professional Summary", before=3)
    add_paragraph(doc, SUMMARY, size=8.55)
    add_heading(doc, "Bar Readiness")
    add_bullets(doc, READINESS, size=8.15)
    add_heading(doc, "Work Experience")
    for company, title, dates, bullets in EXPERIENCE:
        add_experience(doc, company, title, dates, bullets)
    add_heading(doc, "Best Fit")
    add_bullets(doc, BEST_FIT, size=8.1)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build()
