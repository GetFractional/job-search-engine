#!/usr/bin/env python3
"""Build the My Way Ahead Company Operating System DOCX from canonical sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/career-platform/company-os/my-way-ahead-company-operating-system-2026-07-20.md"
PROMPT = ROOT / "docs/career-platform/company-os/new-task-activation-prompt-2026-07-20.txt"
ORG_IMAGE = ROOT / "output/visuals/my-way-ahead/my-way-ahead-company-org-chart-2026-07-19.png"
VISUAL_IMAGE = ROOT / "output/visuals/my-way-ahead/selected-executive-evidence-timeline.png"
OUTPUT = ROOT / "output/doc/My_Way_Ahead_Company_Operating_System_and_Transfer_Package_2026-07-20.docx"

NAVY = "10212D"
INK = "1E303A"
MINT = "49C9A7"
PALE_MINT = "E9F8F3"
SAND = "F7F3EA"
GRAY = "5F6F78"
LINE = "C9D3D7"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=85, bottom=85, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.02
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.2 if not header else 7.4)
        run.font.color.rgb = RGBColor.from_string(WHITE if header else INK)
        if header:
            run.bold = True
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "168A73")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("245B4E")
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.25)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4.2)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Title", 34, NAVY, 0, 10),
        ("Heading 1", 19, NAVY, 16, 6),
        ("Heading 2", 13, "168A73", 11, 4),
        ("Heading 3", 10.5, NAVY, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.font.name = "Aptos"
        callout.font.size = Pt(10)
        callout.font.color.rgb = RGBColor.from_string(NAVY)
        callout.paragraph_format.space_before = Pt(5)
        callout.paragraph_format.space_after = Pt(5)
        callout.paragraph_format.left_indent = Inches(0.12)
        callout.paragraph_format.right_indent = Inches(0.12)

    header = section.header
    p = header.paragraphs[0]
    p.text = "MY WAY AHEAD   |   COMPANY OPERATING SYSTEM"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("168A73")
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), MINT)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.add_run("BOARD PACKET  |  VERSION 1.0  |  2026-07-20")
    for run in fp.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.3)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    fp.add_run(" " * 8)
    add_page_number(fp)


def add_cover(document: Document) -> None:
    label = document.add_paragraph()
    label.paragraph_format.space_before = Pt(72)
    run = label.add_run("BOARD OPERATING PACKET")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("168A73")

    title = document.add_paragraph(style="Title")
    title.add_run("My Way Ahead\nCompany Operating System")
    title.paragraph_format.space_before = Pt(8)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(4)
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("Business plan, objectives, evolving organization, accountability, automation, and clean-task transfer")
    run.font.name = "Aptos Display"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    labels = [
        ("BOARD AND OWNER", "Matt Dimock"),
        ("ACCOUNTABLE CEO", "Codex root, within delegated authority"),
        ("CURRENT STAGE", "Founder candidate, before private external alpha"),
        ("OPERATING NORTH STAR", "Trusted career decisions that change, confirm, or prevent a consequential action"),
    ]
    for row, (key, value) in zip(table.rows, labels):
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(5.2)
        shade(row.cells[0], NAVY)
        shade(row.cells[1], SAND)
        set_cell_text(row.cells[0], key, header=True)
        set_cell_text(row.cells[1], value)
        for run in row.cells[1].paragraphs[0].runs:
            run.font.size = Pt(9)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("PROVISIONAL COMPANY NAME")
    run.font.size = Pt(8)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("168A73")
    p.add_run("\nMy Way Ahead is not legally cleared, purchased, reserved, or authorized for public launch.")

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    run = p.add_run("THE DECISION")
    run.font.size = Pt(8)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("168A73")
    text = (
        "\nBuild the owned product in thin, trustworthy slices. Keep only two active initiatives: "
        "private-alpha readiness and Matt's next-best-job case. Earn external testing, payment, "
        "recurring billing, and scale through evidence rather than feature volume."
    )
    p.add_run(text)
    document.add_page_break()


def add_callout(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE_MINT)
    set_cell_margins(cell, top=150, start=160, bottom=150, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Callout"]
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(NAVY)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(columns):
            cell = table.cell(r_idx, c_idx)
            value = row[c_idx] if c_idx < len(row) else ""
            if r_idx == 0:
                shade(cell, NAVY)
            elif r_idx % 2 == 0:
                shade(cell, "F5F8F7")
            set_cell_text(cell, value, header=(r_idx == 0))
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    idx = 0
    first_h1_skipped = False
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line or line.startswith("<!--"):
            idx += 1
            continue
        if line.startswith("# "):
            if not first_h1_skipped:
                first_h1_skipped = True
            else:
                document.add_heading(line[2:].strip(), level=1)
            idx += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            idx += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            idx += 1
            continue
        if line.startswith("#### "):
            document.add_heading(line[5:].strip(), level=3)
            idx += 1
            continue
        if line.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(document, rows)
            continue
        if line.startswith(">"):
            quote_lines = []
            while idx < len(lines) and lines[idx].strip().startswith(">"):
                quote_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(document, " ".join(quote_lines))
            continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if ordered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            number_run = paragraph.add_run(f"{ordered.group(1)}.  ")
            number_run.bold = True
            number_run.font.color.rgb = RGBColor.from_string("168A73")
            add_inline(paragraph, ordered.group(2))
            idx += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:].strip())
            idx += 1
            continue
        if re.match(r"^\*\*[^*]+:\*\*", line):
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
            idx += 1
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, line)
        idx += 1


def add_image_page(document: Document, title: str, image_path: Path, caption: str, width: float) -> None:
    document.add_page_break()
    document.add_heading(title, level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_prompt_appendix(document: Document, prompt_text: str) -> None:
    document.add_page_break()
    document.add_heading("Copy-and-paste New Task Activation Prompt", level=1)
    p = document.add_paragraph(
        "Use the plain-text companion file for copying. This appendix keeps the complete transfer instructions inside the Board packet."
    )
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(GRAY)
    for block in prompt_text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.endswith(":") and len(block) < 80:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("168A73")
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.08)
            paragraph.paragraph_format.right_indent = Inches(0.08)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, block)
            for run in paragraph.runs:
                run.font.size = Pt(7.9)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_markdown(document, SOURCE.read_text(encoding="utf-8"))
    add_image_page(
        document,
        "Visual Appendix: Actual Company Organization",
        ORG_IMAGE,
        "Permanent departments and real accountability roles. Seats activate, combine, or dissolve as the business stage changes.",
        5.45,
    )
    add_image_page(
        document,
        "Visual Appendix: Selected Product Direction",
        VISUAL_IMAGE,
        "Executive Evidence with the connected Integrity Preview timeline is the selected design authority. Current-build browser comparison is still required.",
        4.05,
    )
    add_prompt_appendix(document, PROMPT.read_text(encoding="utf-8"))
    document.core_properties.title = "My Way Ahead Company Operating System and Transfer Package"
    document.core_properties.subject = "Board operating packet"
    document.core_properties.author = "My Way Ahead CEO Office"
    document.core_properties.comments = "Generated from canonical local sources."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
