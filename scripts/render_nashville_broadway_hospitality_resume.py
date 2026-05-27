from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path("/Users/mattdimock/Documents/Jobs/Job Search")
OUT_DIR = ROOT / "applications" / "hospitality-lanes"
OUT_DOCX = OUT_DIR / "Matt Dimock - Nashville Broadway Hospitality - Resume.docx"
OUT_PDF = OUT_DIR / "Matt Dimock - Nashville Broadway Hospitality - Resume.pdf"

INK = RGBColor(17, 24, 39)
MUTED = RGBColor(75, 85, 99)
RULE = RGBColor(209, 213, 219)
PANEL = RGBColor(247, 247, 245)

PDF_INK = colors.HexColor("#111827")
PDF_MUTED = colors.HexColor("#4B5563")
PDF_RULE = colors.HexColor("#D1D5DB")
PDF_PANEL = colors.HexColor("#F7F7F5")

HEADER = {
    "name": "Matt Dimock",
    "target": "Server | Bartender | Barback",
    "contact": "Mount Juliet, TN  |  805-620-2826  |  mattdim805@gmail.com  |  linkedin.com/in/mattdimock/",
}

SUMMARY = (
    "High-ownership, customer-facing operator moving into Nashville hospitality after 19 years across demanding customers, sales, "
    "operations, retail, training, and service-driven environments. Best fit: server, barback, or bartender roles where maturity, "
    "fast recall, stamina, calm under pressure, and team-first execution matter. Open availability, nights and weekends, reliable "
    "transportation, and Tennessee ABC Server Permit completed."
)

HIGHLIGHT = (
    "Brings owner-level judgment, high-volume customer experience, strong memory, physical stamina, process discipline, and no-ego "
    "support for whatever keeps service moving."
)

SKILLS = [
    ("Guest Service & Sales", "Warm presence; demanding customer communication; recovery judgment; upsell awareness; repeat-guest mindset; polished follow-through."),
    ("Bar & Floor Readiness", "Tennessee ABC Server Permit completed; barback support; server readiness; stocking; restocking; opening and closing routines; inventory counts; clean handoffs."),
    ("Pace, Memory & Detail", "Fast menu learning; recipe/process recall; multitasking; checklists; accuracy under pressure; strong attention to small details."),
    ("Reliability & Stamina", "Open availability; nights and weekends; physically fit; 10K+ daily steps; gym-active; heavy-lifting comfort; steady energy."),
    ("Ownership & Team Fit", "No ego; takes direction; trains quickly; solves problems; protects the guest experience; understands operator priorities."),
]

EXPERIENCE = [
    (
        "Get Fractional",
        "Growth Systems Consultant / Owner",
        "Oct 2022 - Present",
        [
            "Own growth and operating-system work for ecommerce and B2B clients, turning messy priorities into campaign calendars, dashboards, vendor coordination, site improvements, and accountable follow-through.",
            "For OC Ramps, manage Shopify growth execution across SEO, email, onsite conversion, reporting, and agency coordination; helped deliver a record December 2025 month and a 39% H2 vs H1 monthly revenue lift after re-engaging campaign management.",
        ],
    ),
    (
        "Prosper Wireless",
        "Director of Growth & Retention",
        "Sept 2023 - Nov 2025",
        [
            "Built training, onboarding, lifecycle, and support systems for a high-volume telecom operation serving roughly 600K+ activations and customers, reducing onboarding burden by about 50%.",
            "Operated as a right-hand systems leader under shifting program, compliance, partner, and customer demands, keeping execution organized when speed and detail both mattered.",
        ],
    ),
    (
        "Dealer Acceleration Group",
        "Co-Founder / Chief Revenue Officer",
        "Jan 2022 - Dec 2022",
        [
            "Co-founded a fast-moving startup, helped raise $300K in seed funding, and directed 3 major pivots in 12 months based on market signal, customer conversations, and unit economics.",
        ],
    ),
    (
        "Affordable Insurance Quotes",
        "Chief Marketing Officer",
        "Jun 2020 - Dec 2021",
        [
            "Helped scale a customer-facing insurance business from about $290K to $2M in 18 months by improving follow-up, sales systems, affiliate operations, and day-to-day operating structure.",
            "Generated 30K+ leads at a 55% conversion rate and helped build the operating framework that expanded the business from 1 to 16 states.",
        ],
    ),
    (
        "Breakthrough Academy",
        "Marketing Operations Manager",
        "Aug 2019 - Jun 2021",
        [
            "Coordinated events, webinars, partners, vendors, and internal teams in a service-driven membership business where timing, presentation, and responsiveness shaped the customer experience.",
            "Built partner-led webinar and event systems that achieved 30%+ conversion rates and outperformed the prior top event channel.",
        ],
    ),
    (
        "SkyFineUSA",
        "Chief Marketing Officer / Head of Growth",
        "Jan 2019 - Aug 2019",
        [
            "Led ecommerce and operational execution in a fast-paced retail business, coordinating priorities across sales, product, reporting, QA, billing, vendor servicing, inventory, and shipping.",
        ],
    ),
    (
        "Bob's Watches",
        "Director of Marketing",
        "Jan 2017 - Dec 2018",
        [
            "Helped grow luxury ecommerce revenue from $20.5M to $45M in 24 months while leading a 9-person team and 15+ vendors in a business where trust, timing, accuracy, and presentation were revenue-critical.",
        ],
    ),
    (
        "Swell Marketing",
        "SEO Manager / SEO Director",
        "Feb 2016 - Jun 2016",
        [
            "Earned promotion to SEO Director within three months while improving SOPs, onboarding, team execution, and client communication in a demanding agency environment.",
        ],
    ),
    (
        "iMarket Solutions",
        "SEO Manager",
        "May 2013 - Jan 2016",
        [
            "Recruited, trained, and managed a five-person team while overseeing 400+ client accounts at the team level and personally managing 45 VIP client relationships.",
        ],
    ),
    (
        "National Positions",
        "SEO Specialist / Director of Special Projects",
        "Sep 2007 - May 2013",
        [
            "Managed more than 120 client accounts per month earlier in tenure, building durable habits around pace, recall, issue tracking, customer communication, and follow-through.",
        ],
    ),
    (
        "Lowe's Home Improvement",
        "Team Leader, Paint Department",
        "Mar 2005 - Feb 2007",
        [
            "Promoted into a team-lead role within six months and helped run a four-person department through customer service, stocking, merchandising, inventory counts, opening, closing, and hands-on floor execution.",
        ],
    ),
    (
        "Boething Treeland Nursery",
        "Guide",
        "Dec 2003 - Aug 2004",
        [
            "Promoted from Loader to Guide within three months, supporting customers, loading, restocking, grounds upkeep, and physically demanding daily operations in a fast-paced retail environment.",
        ],
    ),
]


def set_doc_margins(section):
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.54)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)


def set_font(run, size, color=INK, bold=False, name="Aptos", italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_bottom_border(paragraph, color="D1D5DB", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def shade_cell(cell, color="F7F7F5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_doc_header(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(HEADER["name"]), 23, INK, True, "Aptos Display")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(HEADER["target"]), 10, MUTED, True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(HEADER["contact"]), 8.3, MUTED)
    set_bottom_border(p)


def add_doc_heading(doc, text, before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(text.upper()), 8.5, INK, True)


def add_doc_body(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    set_font(p.add_run(text), 8.9, INK, bold)


def add_doc_skills(doc):
    table = doc.add_table(rows=len(SKILLS), cols=2)
    table.autofit = True
    for row, (label, body) in zip(table.rows, SKILLS):
        for cell in row.cells:
            shade_cell(cell)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 8.4, INK, True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(body), 8.15, INK)


def add_doc_experience(doc, company, title, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(1.5)
    set_font(p.add_run(company), 9.5, INK, True)
    set_font(p.add_run(f"  |  {title}  |  {dates}"), 8.2, MUTED, italic=True)
    for bullet in bullets:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(11)
        p.paragraph_format.first_line_indent = Pt(-8)
        p.paragraph_format.line_spacing = 1.04
        set_font(p.add_run(f"- {bullet}"), 8.15, INK)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    set_doc_margins(section)
    add_doc_header(doc)
    add_doc_heading(doc, "Summary", before=0)
    add_doc_body(doc, SUMMARY)
    add_doc_body(doc, HIGHLIGHT, bold=True)
    add_doc_heading(doc, "Skills", before=8)
    add_doc_skills(doc)
    add_doc_heading(doc, "Experience", before=10)
    for company, title, dates, bullets in EXPERIENCE:
        add_doc_experience(doc, company, title, dates, bullets)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


def pdf_bullets(items):
    return "<br/>".join(f"- {item}" for item in items)


def add_pdf_header(story, styles):
    story.extend(
        [
            Paragraph(HEADER["name"], styles["Name"]),
            Paragraph(HEADER["target"], styles["Target"]),
            Paragraph(HEADER["contact"], styles["Contact"]),
            Table([[""]], colWidths=[7.15 * inch], rowHeights=[0.75], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), PDF_RULE)])),
            Spacer(1, 11),
        ]
    )


def add_pdf_experience(story, styles, entries):
    for company, title, dates, bullets in entries:
        story.append(Paragraph(company, styles["Company"]))
        story.append(Paragraph(f"{title}  |  {dates}", styles["Meta"]))
        if bullets:
            story.append(Paragraph(pdf_bullets(bullets), styles["Bullet"]))


def build_pdf():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=letter,
        rightMargin=0.58 * inch,
        leftMargin=0.58 * inch,
        topMargin=0.52 * inch,
        bottomMargin=0.48 * inch,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("Name", fontName="Helvetica-Bold", fontSize=22.5, leading=23.5, textColor=PDF_INK, spaceAfter=1))
    styles.add(ParagraphStyle("Target", fontName="Helvetica-Bold", fontSize=9.3, leading=10.5, textColor=PDF_MUTED, spaceAfter=1))
    styles.add(ParagraphStyle("Contact", fontName="Helvetica", fontSize=8.0, leading=9.0, textColor=PDF_MUTED, spaceAfter=5))
    styles.add(ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=8.1, leading=9.0, textColor=PDF_INK, spaceBefore=9, spaceAfter=4))
    styles.add(ParagraphStyle("Body", fontName="Helvetica", fontSize=8.2, leading=9.5, textColor=PDF_INK, spaceAfter=5))
    styles.add(ParagraphStyle("BodyBold", fontName="Helvetica-Bold", fontSize=8.05, leading=9.2, textColor=PDF_INK, spaceAfter=5))
    styles.add(ParagraphStyle("SkillLabel", fontName="Helvetica-Bold", fontSize=7.35, leading=8.2, textColor=PDF_INK, spaceAfter=0))
    styles.add(ParagraphStyle("SkillBody", fontName="Helvetica", fontSize=7.25, leading=8.25, textColor=PDF_INK, spaceAfter=0))
    styles.add(ParagraphStyle("Company", fontName="Helvetica-Bold", fontSize=8.8, leading=9.6, textColor=PDF_INK, spaceBefore=7.0, spaceAfter=0.4))
    styles.add(ParagraphStyle("Meta", fontName="Helvetica-Oblique", fontSize=7.55, leading=8.3, textColor=PDF_MUTED, spaceAfter=0.4))
    styles.add(ParagraphStyle("ResumeBullet", fontName="Helvetica", fontSize=7.55, leading=8.65, textColor=PDF_INK, leftIndent=8, firstLineIndent=-6.5, spaceAfter=0.7))

    skill_rows = []
    for label, body in SKILLS:
        skill_rows.append([Paragraph(label, styles["SkillLabel"]), Paragraph(body, styles["SkillBody"])])
    skill_table = Table(skill_rows, colWidths=[1.55 * inch, 5.49 * inch])
    skill_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_PANEL),
                ("BOX", (0, 0), (-1, -1), 0.35, PDF_RULE),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, PDF_RULE),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 3.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )

    story = []
    add_pdf_header(story, styles)
    story.append(Paragraph("SUMMARY", styles["Section"]))
    story.append(Paragraph(SUMMARY, styles["Body"]))
    story.append(Paragraph(HIGHLIGHT, styles["BodyBold"]))
    story.append(Paragraph("SKILLS", styles["Section"]))
    story.append(skill_table)
    story.append(Paragraph("EXPERIENCE", styles["Section"]))
    add_pdf_experience(story, styles, EXPERIENCE[:5])
    story.append(PageBreak())
    add_pdf_header(story, styles)
    story.append(Paragraph("EXPERIENCE", styles["Section"]))
    add_pdf_experience(story, styles, EXPERIENCE[5:])

    doc.build(story)


def main():
    build_docx()
    build_pdf()
    print(OUT_DOCX)
    print(OUT_PDF)


if __name__ == "__main__":
    main()
